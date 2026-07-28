"""
Parses the legacy 'CCC Accounts' workbook into structured match records.

The two Entry sheets use different block shapes, and neither is reliably
identified by its labels (a payment row's label may be a type like
'Practice match', or just another venue name). The one invariant that holds
across both is that only the payment row carries the Collected total in
column C, so detection anchors on that.

  'Entry' (Jul 2022 - Nov 2023), 2-row blocks, flat fee, payments only:
    payment  date | venue | C=collected | per-player payments | expenses
    mode          | 'CCC League - Mnn'  | per-player payment modes

  'Entry - New' (Nov 2023 onward), 3-row blocks, explicit per-player charges:
    charge   date | venue        | negative per-player match fee
    payment  date | type/venue   | C=collected | per-player payments | expenses
    mode     date | 'Match - N' or type | per-player payment modes

Rows carrying money that belong to no fixture (opening balances, kit
settlements passing through the ledger) are returned separately rather than
being folded into a match.
"""
import openpyxl, re, json, sys, datetime

SRC = sys.argv[1] if len(sys.argv) > 1 else \
    "/root/.claude/uploads/70d8d0d5-9bf3-5377-a487-a89dd37b7cf2/b6a588db-CCC__Accounts.xlsx"

OVERS_RE = re.compile(r'(\d+)\s*over', re.I)
NUM_RE = re.compile(r'(\d+)\s*$')
LEAGUE_RE = re.compile(r'league', re.I)
PRACTICE_RE = re.compile(r'practice', re.I)
CATEGORIES = {"ball fee", "ground fee", "mom", "league", "water", "bakery",
              "food", "fuel", "umpire fee", "others"}
# Column headers that are neither players nor expense categories.
NOT_A_PLAYER = {"present", "balance", "cash", "g-pay", "gpay", "upi", "adjust",
                "pending", "absent", "total", "ball", "car", "guest", "date",
                "particulars", "collected"}


def columns(ws):
    players, cats = {}, {}
    for c in range(5, ws.max_column + 1):
        v = ws.cell(1, c).value
        if not isinstance(v, str) or not v.strip():
            continue
        v = v.strip()
        if v.lower() in CATEGORIES:
            cats[c] = v
        elif not cats and v.lower() not in NOT_A_PLAYER:
            players[c] = v
    return players, cats


def has_text(ws, r, cols):
    return any(isinstance(ws.cell(r, c).value, str) and ws.cell(r, c).value.strip()
               for c in cols)


def has_negative(ws, r, cols):
    return any(isinstance(ws.cell(r, c).value, (int, float)) and ws.cell(r, c).value < 0
               for c in cols)


def parse_sheet(ws):
    players, cats = columns(ws)
    matches, used = [], set()

    for r in range(3, ws.max_row + 1):
        if not isinstance(ws.cell(r, 3).value, (int, float)):
            continue                      # not a payment row
        label = (ws.cell(r, 2).value or "")
        label = label.strip() if isinstance(label, str) else ""
        if label.lower() in ("opening", "balance"):
            continue

        # A charge row sits directly above only when it carries negative
        # per-player amounts for the same fixture (the newer sheet's shape).
        charge_r = r - 1 if r > 3 and has_negative(ws, r - 1, players) \
            and not isinstance(ws.cell(r - 1, 3).value, (int, float)) else None
        mode_r = r + 1 if r + 1 <= ws.max_row and has_text(ws, r + 1, players) else None

        date = ws.cell(r, 1).value or (ws.cell(charge_r, 1).value if charge_r else None)
        venue = (ws.cell(charge_r, 2).value if charge_r else label) or label
        venue = venue.strip() if isinstance(venue, str) else None
        mode_label = str(ws.cell(mode_r, 2).value or "") if mode_r else ""

        # Type and overs may be stated on either the payment row's label or
        # the mode row's; league is the default since practice is the exception.
        blob = f"{label} {mode_label}"
        is_practice = bool(PRACTICE_RE.search(blob))
        overs_m = OVERS_RE.search(blob)
        num_m = NUM_RE.search(mode_label.replace("overs", "").replace("over", ""))

        roster = []
        for c, name in players.items():
            charge = ws.cell(charge_r, c).value if charge_r else None
            paid = ws.cell(r, c).value
            mode = ws.cell(mode_r, c).value if mode_r else None
            charge = -charge if isinstance(charge, (int, float)) else None
            paid = paid if isinstance(paid, (int, float)) else 0
            mode = mode.strip() if isinstance(mode, str) and mode.strip() else None
            if charge in (None, 0) and paid == 0 and mode is None:
                continue
            roster.append({"player": name, "payable": charge, "collected": paid,
                           "mode": mode})

        # Expenses are stored negative (money out). Flip the sign so the app
        # holds positive amounts, and let genuinely positive cells become
        # negative expenses — those are refunds and reimbursements.
        expenses = [{"category": name, "amount": -ws.cell(r, c).value,
                     "note": (ws.cell(mode_r, c).value.strip()
                              if mode_r and isinstance(ws.cell(mode_r, c).value, str) else None)}
                    for c, name in cats.items()
                    if isinstance(ws.cell(r, c).value, (int, float)) and ws.cell(r, c).value]

        matches.append({
            "date": date.date().isoformat() if isinstance(date, datetime.datetime) else None,
            "venue": venue,
            "isLeague": not is_practice,
            "overs": int(overs_m.group(1)) if overs_m else None,
            "matchNumber": int(num_m.group(1)) if num_m else None,
            "collected": ws.cell(r, 3).value,
            "roster": roster,
            "expenses": expenses,
        })
        used.update({r} | ({charge_r} if charge_r else set()) | ({mode_r} if mode_r else set()))

    leftovers = []
    for r in range(3, ws.max_row + 1):
        if r in used:
            continue
        net = sum(ws.cell(r, c).value for c in list(players) + list(cats)
                  if isinstance(ws.cell(r, c).value, (int, float)))
        if net:
            leftovers.append({"row": r, "label": ws.cell(r, 2).value, "net": net})
    return players, cats, matches, leftovers


if __name__ == "__main__":
    wb = openpyxl.load_workbook(SRC, data_only=True)
    out = {}
    for sheet in ("Entry", "Entry - New"):
        ws = wb[sheet]
        players, cats, matches, leftovers = parse_sheet(ws)
        out[sheet] = {"matches": matches, "leftovers": leftovers,
                      "players": list(players.values())}
        # Checksum: money the parser claims vs the column totals it came from.
        claimed = {}
        for m in matches:
            for x in m["roster"]:
                claimed[x["player"]] = claimed.get(x["player"], 0) \
                    - (x["payable"] or 0) + x["collected"]
        drift = {}
        for c, name in players.items():
            col = sum(ws.cell(r, c).value for r in range(3, ws.max_row + 1)
                      if isinstance(ws.cell(r, c).value, (int, float)))
            d = round(col - claimed.get(name, 0), 2)
            if d:
                drift[name] = d
        cat_drift = {}
        for c, name in cats.items():
            col = -sum(ws.cell(r, c).value for r in range(3, ws.max_row + 1)
                       if isinstance(ws.cell(r, c).value, (int, float)))
            got = sum(e["amount"] for m in matches for e in m["expenses"] if e["category"] == name)
            if round(col - got, 2):
                cat_drift[name] = round(col - got, 2)
        print(f"{sheet}: {len(matches)} matches | unclaimed rows {len(leftovers)} "
              f"| player drift {sum(abs(v) for v in drift.values()):.0f} "
              f"| category drift {sum(abs(v) for v in cat_drift.values()):.0f}", file=sys.stderr)
        if leftovers:
            print(f"    unclaimed: {[(l['label'], l['net']) for l in leftovers]}", file=sys.stderr)
        if cat_drift:
            print(f"    category drift detail: {cat_drift}", file=sys.stderr)
    json.dump(out, open("/tmp/parsed.json", "w"), indent=1)
